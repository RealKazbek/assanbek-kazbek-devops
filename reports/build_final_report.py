"""Build the final DevOps laboratory report as matching DOCX and PDF files.

The report intentionally distinguishes genuine text/API evidence from screenshots that
could not be captured under the command-line-only constraint.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
DOCX_PATH = REPORTS / "Assanbek_Kazbek_DevOps_Lab_Report.docx"
PDF_PATH = REPORTS / "Assanbek_Kazbek_DevOps_Lab_Report.pdf"
FONT_DIR = Path("/System/Library/Fonts/Supplemental")

STUDENT = {
    "Name": "Kazbek",
    "Surname": "Assanbek",
    "Group": "IT2-2302",
    "Student ID": "37765",
    "Course": "DevOps",
}


TASKS = [
    ("Task 1 - Project Environment", "Create the personalized project structure, student information file, configuration and log files. Demonstrate pwd, ls, cd, mkdir, touch, cp, mv, rm, cat, head and tail.", "The project contains app, scripts, config, logs and backup. Assanbek_Kazbek_info.txt contains the required five student fields. The Linux session demonstrates every required command and docs/OPERATIONS.md explains each one.", "evidence/linux/linux_commands.txt"),
    ("Task 2 - Search and File Processing", "Create application.log with at least ten messages; use find, grep ERROR, grep WARNING, an ERROR count and tail -n 5.", "logs/application.log contains realistic messages. The recorded Linux session finds the file, finds ERROR and WARNING lines, counts two ERROR lines and prints the last five entries.", "logs/application.log; evidence/linux/linux_commands.txt"),
    ("Task 3 - File Permissions", "Create private.txt and public.txt with different permissions. Use chmod and ls -l, then explain r, w, x, owner, group and others.", "private.txt was set to 600 and public.txt to 644. The Linux output shows -rw------- and -rw-r--r--; the operational notes explain why the files differ.", "private.txt; public.txt; docs/OPERATIONS.md; evidence/linux/linux_commands.txt"),
    ("Task 4 - Bash Script", "Create an executable personalized system script. It must show student and system data using whoami, hostname, date, uname, df and free, and check the information file.", "scripts/Assanbek_Kazbek_system.sh is executable. Its genuine Linux run prints all requested fields, disk and memory data, and confirms that the information file exists.", "scripts/Assanbek_Kazbek_system.sh; evidence/linux/linux_commands.txt"),
    ("Task 5 - Processes", "Demonstrate ps, ps aux and top. Identify PID, owner, CPU and memory, and explain process management and the difference between ps and top.", "The Linux record includes ps, ps aux and top output. ps aux shows PID, USER, %CPU and %MEM. docs/OPERATIONS.md gives the required explanation.", "docs/OPERATIONS.md; evidence/linux/linux_commands.txt"),
    ("Task 6 - Git Repository", "Initialize Git, check user.name and user.email, and show git status.", "The repository has a real initial commit and current Git identity is Kazbek / 37765@iitu.edu.kz. The repository was clean before this report update.", "local .git metadata; evidence/github/git_history_latest.txt"),
    ("Task 7 - Meaningful Commit History", "Create at least five meaningful commits and show git log and git log --oneline.", "The graph includes project initialization, application, Linux verification, Docker verification, Compose verification, documentation, Jenkins configuration and final verification commits. These are substantive changes rather than empty commits.", "evidence/github/git_history_latest.txt"),
    ("Task 8 - Gitignore", "Create .gitignore for *.log, backup/ and .env. Show ignored files and explain why the file is useful.", "The required three patterns are present. git check-ignore -v recorded matches for logs/application.log, backup and .env. The explanation notes that logs, backups and local secrets should not enter history.", ".gitignore; evidence/gitignore_check.txt"),
    ("Task 9 - Branching and Merge", "Create development, make a meaningful change, commit it, merge it into main, show the branch/graph and explain branches.", "development contains documentation commit e602227 and was merged by e5b03c7. The graph is preserved in the latest Git evidence.", "docs/OPERATIONS.md; evidence/github/git_history_latest.txt"),
    ("Task 10 - GitHub", "Create a repository, connect origin, verify remote, push main and ensure the complete project is on GitHub.", "GitHub CLI verified the public repository RealKazbek/assanbek-kazbek-devops and main branch. Local HEAD and origin/main matched at the time of verification.", "evidence/github/repository_verification.json; evidence/github/git_history_latest.txt"),
    ("Task 11 - Dockerfile", "Create a Dockerfile with FROM, WORKDIR, COPY, RUN, EXPOSE and CMD. Run an application that displays the student information and success message.", "The Python standard-library application produces the requested text. Dockerfile uses python:3.13-alpine, /app, a non-root student user, port 8080 and python -m app.app.", "Dockerfile; app/app.py; tests/test_app.py; evidence/docker/docker_verification.txt"),
    ("Task 12 - Build and Run Docker Container", "Build the personalized image, run the named container, and demonstrate docker images, docker ps and docker ps -a.", "The Docker record proves a successful build of assanbek-kazbek-devops and a running assanbek-kazbek-container. Current Docker inspection also found the container running.", "evidence/docker/docker_verification.txt"),
    ("Task 13 - Environment Variables", "Pass STUDENT_NAME, STUDENT_SURNAME, STUDENT_GROUP and STUDENT_ID to the container. Show changed output and explain why variables are useful and secrets must not be hard-coded.", "app/app.py reads all four values. A real override run printed Changed Student, DEMO-1 and 99999. The operational notes explain runtime configuration and secret handling.", "app/app.py; tests/test_app.py; docs/OPERATIONS.md; evidence/docker/docker_verification.txt"),
    ("Task 14 - Container Management", "Demonstrate docker logs, stop, start, rm and, if applicable, exec. Explain image versus container.", "The evidence shows logs, exec, stop and start for the required container and an isolated removal demo. The notes correctly distinguish an immutable image from a running or stopped container instance.", "evidence/docker/docker_verification.txt; evidence/docker/compose_and_management.txt; docs/OPERATIONS.md"),
    ("Task 15 - Docker Compose", "Create docker-compose.yml, start the application with docker compose up, check it, stop it with docker compose down and explain Compose.", "The compose service supplies all student environment values, uses a non-conflicting container name, and was built, started, logged and removed through compose down.", "docker-compose.yml; evidence/docker/compose_and_management.txt; docs/OPERATIONS.md"),
    ("Task 16 - Jenkins Job", "Create the exact Jenkins Job Assanbek_Kazbek_DevOps using the GitHub repository.", "The preserved job XML uses the public repository URL, */main and Jenkinsfile. The successful console begins Started Assanbek_Kazbek_DevOps #2.", "jenkins/job-config.xml; evidence/jenkins/console_build_2.txt"),
    ("Task 17 - Jenkinsfile", "Create a declarative Jenkinsfile with meaningful Checkout, Build, Test, Docker Build and Docker Run stages.", "Jenkinsfile defines all five stages. The stages check out SCM, compile Python, run two unit tests, build the image, run the container and inspect output.", "Jenkinsfile; evidence/jenkins/console_build_2.txt"),
    ("Task 18 - Jenkins and GitHub", "Configure Jenkins to check out the GitHub project, build it, test it, build a Docker image and run a Docker container.", "Build #2 fetched https://github.com/RealKazbek/assanbek-kazbek-devops.git, checked out e16e061, then executed all required stages.", "jenkins/job-config.xml; evidence/jenkins/console_build_2.txt"),
    ("Task 19 - Jenkins Environment Variables", "Configure all four student variables and display Student, Group and Student ID in Jenkins Console Output.", "The Jenkins environment block contains the four values. Console output contains Student: Kazbek Assanbek, Group: IT2-2302 and Student ID: 37765.", "Jenkinsfile; evidence/jenkins/console_build_2.txt"),
    ("Task 20 - Successful Pipeline", "Run the pipeline successfully and demonstrate Dashboard, job name, pipeline stages, console output, Docker image/container and the Git commit used by Jenkins.", "Build #2 executed all five stages and ended Finished: SUCCESS. It built assanbek-kazbek-devops:jenkins, ran assanbek-kazbek-jenkins-app and checked out e16e061. Text/API evidence is genuine; required visual Jenkins screenshots were not captured under the command-line-only constraint.", "evidence/jenkins/console_build_2.txt; evidence/jenkins/console_build_2_api.txt; evidence/jenkins/controller_and_docker_verification.txt; evidence/SCREENSHOT_CAPTURE_STATUS.md"),
    ("Task 21 - Complete DevOps Workflow", "Demonstrate Linux to Git to GitHub to Jenkins to build, test, Docker build, Docker run and SUCCESS. Explain each connection and failure handling.", "docs/OPERATIONS.md explains the actual implemented chain, image/container distinction and failure behavior. Linux, GitHub, Docker and Jenkins logs provide genuine technical evidence. The student still must present the workflow in the required own-voice video.", "docs/OPERATIONS.md; README.md; VIDEO_GUIDE.md; evidence/linux/; evidence/github/; evidence/jenkins/"),
]

EVIDENCE_SNIPPETS = [
    ("Linux command and permissions evidence", """Linux environment: Linux ... aarch64 Linux
find . -name application.log
./logs/application.log
-rw------- ... private.txt
-rw-r--r-- ... public.txt
Student information file exists.
ps aux ... USER PID %CPU %MEM"""),
    ("Docker and Compose evidence", """assanbek-kazbek-devops:latest
assanbek-kazbek-container ... Up ... 8081->8080/tcp
Name: Kazbek
Surname: Assanbek
Group: IT2-2302
Student ID: 37765
Application is running successfully!"""),
    ("GitHub and branch evidence", """* f50498d (HEAD -> main, origin/main) Verify successful Jenkins pipeline and update final evidence
*   e5b03c7 Merge development documentation updates
| * e602227 (development) Document Linux and Docker operating procedures
RealKazbek/assanbek-kazbek-devops | PUBLIC | default=main"""),
    ("Jenkins successful pipeline evidence", """Checking out Revision e16e0613786fe1df391589a7f283d7f5af0c9bb1
[Pipeline] { (Checkout)
[Pipeline] { (Build)
[Pipeline] { (Test)
[Pipeline] { (Docker Build)
[Pipeline] { (Docker Run)
Student: Kazbek Assanbek
Group: IT2-2302
Student ID: 37765
Finished: SUCCESS"""),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_docx(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.22
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 14)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    if "Code Evidence" not in styles:
        code = styles.add_style("Code Evidence", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Courier New"
        code.font.size = Pt(9)
        code.font.color.rgb = RGBColor(0, 0, 0)
        code.paragraph_format.space_after = Pt(7)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def docx_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def docx_table(document, rows, widths=None):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        hdr[idx].text = value
        set_cell_shading(hdr[idx], "D9D9D9")
        for run in hdr[idx].paragraphs[0].runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True
    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            for p in cells[idx].paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    document.add_paragraph()
    return table


def build_docx():
    doc = Document()
    configure_docx(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("DevOps Laboratory Work")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subrun = subtitle.add_run("Linux Git Docker and Jenkins")
    subrun.bold = True
    subrun.font.name = "Times New Roman"
    subrun.font.size = Pt(16)
    for _ in range(4):
        doc.add_paragraph()
    for line in ("Student: Kazbek Assanbek", "Group: IT2-2302", "Student ID: 37765", "Course: DevOps", "Report date: 9 October 2026"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line)
    doc.add_page_break()

    docx_heading(doc, "1. Student Information")
    docx_table(doc, [["Field", "Value"]] + [[key, value] for key, value in STUDENT.items()], [5, 11])
    doc.add_paragraph("The personalized project directory is assanbek_kazbek. The repository is RealKazbek/assanbek-kazbek-devops. The Docker image is assanbek-kazbek-devops, the manual Docker container is assanbek-kazbek-container, and the Jenkins Job is Assanbek_Kazbek_DevOps.")
    docx_heading(doc, "2. Introduction and Objective")
    doc.add_paragraph("The objective of this individual laboratory work is to implement and explain one connected DevOps workflow: Linux, Git, GitHub, Docker, Jenkins, testing and deployment of a small personalized application. The project uses a genuine Linux Docker environment for Linux-only commands. The GitHub repository is public. The Jenkins pipeline was actually executed and its console log records Finished: SUCCESS.")
    doc.add_paragraph("This report uses only real implementation files and captured command or API evidence. It does not label command output as a screenshot. Where a required visual screenshot could not be captured without graphical interaction, the limitation is stated clearly.")

    section_titles = {
        1: "3. Linux Tasks 1 to 5",
        6: "4. Git and GitHub Tasks 6 to 10",
        11: "5. Docker Tasks 11 to 15",
        16: "6. Jenkins Tasks 16 to 20",
        21: "7. Integrated DevOps Workflow Task 21",
    }
    for index, (title_text, requirement, implementation, evidence) in enumerate(TASKS, start=1):
        if index in section_titles:
            docx_heading(doc, section_titles[index])
        docx_heading(doc, title_text, level=2)
        doc.add_paragraph("Instructor requirement: " + requirement)
        doc.add_paragraph("Implementation and actual result: " + implementation)
        doc.add_paragraph("Evidence: " + evidence)

    doc.add_page_break()
    docx_heading(doc, "8. Real Evidence Extracts")
    doc.add_paragraph("The following extracts are copied from genuine saved command or Jenkins console records. Full files are available at the listed paths.")
    for title, snippet in EVIDENCE_SNIPPETS:
        docx_heading(doc, title, level=2)
        para = doc.add_paragraph(style="Code Evidence")
        para.add_run(snippet)

    docx_heading(doc, "9. Screenshot Evidence Status")
    doc.add_paragraph("The assignment asks for visual demonstration of GitHub and Jenkins screens. No browser, desktop, mouse or keyboard control was permitted while this report was prepared, so no genuine screenshots could be captured. No mockups or generated images are included. The exact blocked list and safe manual capture instructions are in evidence/SCREENSHOT_CAPTURE_STATUS.md.")
    doc.add_paragraph("Existing genuine evidence for the successful Jenkins build is textual: evidence/jenkins/console_build_2.txt and evidence/jenkins/console_build_2_api.txt show the five stages, checked-out commit e16e061 and Finished: SUCCESS. The current GitHub main commit is f50498d; it was created after the successful build to retain verification material. This difference is documented and does not by itself invalidate the build.")

    docx_heading(doc, "10. Conclusion")
    doc.add_paragraph("The technical project satisfies the Linux, Git, GitHub, Docker, Docker Compose and Jenkins implementation requirements with real command evidence. The Jenkins build successfully checked out the public repository, built and tested the Python application, built a Docker image and ran a Docker container. The main remaining submission action is the student's own continuous voice video and the capture of genuine graphical screenshots during that demonstration.")
    doc.add_paragraph("Before recording, the student should review VIDEO_GUIDE.md, use the real project and public repository, and never replace missing evidence with generated or reconstructed images.")
    docx_heading(doc, "11. Final Verification Summary")
    docx_table(doc, [
        ["Area", "Result", "Evidence"],
        ["Linux Tasks 1 to 5", "Verified by genuine Linux command output", "evidence/linux/linux_commands.txt"],
        ["Git and GitHub Tasks 6 to 10", "Verified by repository graph, origin and GitHub CLI", "evidence/github/"],
        ["Docker Tasks 11 to 15", "Verified by build, container and Compose records", "evidence/docker/"],
        ["Jenkins Tasks 16 to 19", "Verified by job XML and successful console execution", "jenkins/job-config.xml; evidence/jenkins/"],
        ["Task 20 visual evidence", "Partial: textual success is verified; screenshots are pending", "evidence/SCREENSHOT_CAPTURE_STATUS.md"],
    ], [4.2, 6.3, 5.4])
    docx_heading(doc, "12. Manual Submission Actions", level=1)
    doc.add_paragraph("The student must record one continuous 10 to 20 minute screen video with their own voice. During the recording, show the GitHub main branch and merge, Linux commands, Docker evidence, Compose, the Jenkins Job, pipeline stages, console output, Docker image/container and Finished: SUCCESS. Save only genuine screenshots taken from the displayed real interfaces.")
    doc.save(DOCX_PATH)


def register_pdf_fonts():
    pdfmetrics.registerFont(TTFont("TNR", str(FONT_DIR / "Times New Roman.ttf")))
    pdfmetrics.registerFont(TTFont("TNR-Bold", str(FONT_DIR / "Times New Roman Bold.ttf")))


def pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(colors.black)
    canvas.setFont("TNR", 10)
    canvas.drawCentredString(A4[0] / 2, 1.05 * cm, f"DevOps Laboratory Work - Kazbek Assanbek - Page {doc.page}")
    canvas.restoreState()


def build_pdf():
    register_pdf_fonts()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName="TNR-Bold", fontSize=22, leading=27, alignment=TA_CENTER, textColor=colors.black, spaceAfter=14)
    subtitle = ParagraphStyle("ReportSubtitle", parent=styles["BodyText"], fontName="TNR-Bold", fontSize=16, leading=20, alignment=TA_CENTER, textColor=colors.black, spaceAfter=9)
    body = ParagraphStyle("ReportBody", parent=styles["BodyText"], fontName="TNR", fontSize=14, leading=18, alignment=TA_LEFT, textColor=colors.black, spaceAfter=8)
    heading = ParagraphStyle("ReportHeading", parent=styles["Heading1"], fontName="TNR-Bold", fontSize=16, leading=20, textColor=colors.black, spaceBefore=10, spaceAfter=7)
    heading2 = ParagraphStyle("ReportHeading2", parent=styles["Heading2"], fontName="TNR-Bold", fontSize=14, leading=17, textColor=colors.black, spaceBefore=8, spaceAfter=5)
    code = ParagraphStyle("ReportCode", parent=styles["Code"], fontName="Courier", fontSize=8.5, leading=10.5, textColor=colors.black, backColor=colors.whitesmoke, borderColor=colors.lightgrey, borderWidth=0.3, borderPadding=5, spaceAfter=9)
    table_text = ParagraphStyle("ReportTable", parent=body, fontName="TNR", fontSize=9.5, leading=11.5, spaceAfter=0)

    def p(text, style=body):
        return Paragraph(text.replace("&", "&amp;"), style)

    def make_table(rows, widths):
        converted = [[p(cell, table_text) for cell in row] for row in rows]
        tbl = Table(converted, colWidths=widths, repeatRows=1, hAlign="LEFT")
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
            ("FONTNAME", (0, 0), (-1, 0), "TNR-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#A6A6A6")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return tbl

    story = [Spacer(1, 5.0 * cm), p("DevOps Laboratory Work", title), p("Linux Git Docker and Jenkins", subtitle), Spacer(1, 1.2 * cm)]
    for line in ("Student: Kazbek Assanbek", "Group: IT2-2302", "Student ID: 37765", "Course: DevOps", "Report date: 9 October 2026"):
        story.append(p(line, ParagraphStyle("cover" + line, parent=body, alignment=TA_CENTER, spaceAfter=7)))
    story += [Spacer(1, 4.0 * cm), p("This report documents the real implementation and verified evidence for an individual DevOps laboratory project. It distinguishes command/API records from missing visual screenshots."), PageBreak()]

    story += [p("1. Student Information", heading), make_table([["Field", "Value"]] + [[key, value] for key, value in STUDENT.items()], [5 * cm, 10.3 * cm]), p("Personalized names: project directory assanbek_kazbek; repository RealKazbek/assanbek-kazbek-devops; image assanbek-kazbek-devops; container assanbek-kazbek-container; Jenkins Job Assanbek_Kazbek_DevOps."), p("2. Introduction and Objective", heading), p("This individual laboratory implements one connected workflow: Linux, Git, GitHub, Docker, Jenkins, testing and Docker deployment. Linux-only commands were executed in a genuine Linux Docker environment. The GitHub repository is public and the Jenkins console records a real successful run."), p("This report does not invent screenshots. It cites genuine command and API evidence and explicitly identifies visual evidence that still must be captured manually." )]

    section_titles = {1: "3. Linux Tasks 1 to 5", 6: "4. Git and GitHub Tasks 6 to 10", 11: "5. Docker Tasks 11 to 15", 16: "6. Jenkins Tasks 16 to 20", 21: "7. Integrated DevOps Workflow Task 21"}
    for index, (title_text, requirement, implementation, evidence) in enumerate(TASKS, start=1):
        if index in section_titles:
            story.append(p(section_titles[index], heading))
        story += [p(title_text, heading2), p("<b>Instructor requirement:</b> " + requirement), p("<b>Implementation and actual result:</b> " + implementation), p("<b>Evidence:</b> " + evidence)]

    story += [PageBreak(), p("8. Real Evidence Extracts", heading), p("The extracts below come from real saved terminal and Jenkins console output. Full files remain in the project evidence directory.")]
    for title_text, snippet in EVIDENCE_SNIPPETS:
        story += [p(title_text, heading2), Preformatted(snippet, code)]
    story += [p("9. Screenshot Evidence Status", heading), p("The assignment requires visual demonstration of GitHub and Jenkins. This report was prepared under a command-line-only restriction: no browser, desktop, mouse, keyboard or graphical Jenkins interface was controlled. Therefore no genuine screenshots were captured and no mockups or generated images are included."), p("The blocked list and manual capture instructions are stored in evidence/SCREENSHOT_CAPTURE_STATUS.md. Genuine textual evidence remains available: evidence/jenkins/console_build_2.txt shows all five stages, commit e16e061 and Finished: SUCCESS. The current GitHub main commit is f50498d because evidence material was committed after the successful build."), p("10. Conclusion", heading), p("The Linux, Git, GitHub, Docker, Compose and Jenkins technical work is implemented and supported by real command evidence. The remaining submission work is the student's own continuous voice video and the capture of genuine graphical screenshots while presenting the existing project."), p("11. Final Verification Summary", heading), make_table([["Area", "Result", "Evidence"], ["Linux Tasks 1 to 5", "Verified by genuine Linux command output", "evidence/linux/linux_commands.txt"], ["Git and GitHub Tasks 6 to 10", "Verified by repository graph, origin and GitHub CLI", "evidence/github/"], ["Docker Tasks 11 to 15", "Verified by build, container and Compose records", "evidence/docker/"], ["Jenkins Tasks 16 to 19", "Verified by job XML and successful console execution", "jenkins/job-config.xml; evidence/jenkins/"], ["Task 20 visual evidence", "Partial: textual success is verified; screenshots are pending", "evidence/SCREENSHOT_CAPTURE_STATUS.md"]], [3.5 * cm, 6.0 * cm, 5.8 * cm]), p("12. Manual Submission Actions", heading), p("The student must record one continuous 10 to 20 minute screen video with their own voice. During the recording, show the GitHub main branch and merge, Linux commands, Docker evidence, Compose, the Jenkins Job, pipeline stages, console output, Docker image/container and Finished: SUCCESS. Save only genuine screenshots taken from the displayed real interfaces.")]

    document = SimpleDocTemplate(str(PDF_PATH), pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=2.0 * cm, bottomMargin=2.3 * cm, title="DevOps Laboratory Work", author="Kazbek Assanbek")
    document.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
